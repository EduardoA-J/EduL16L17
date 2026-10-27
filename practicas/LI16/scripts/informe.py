# -*- coding: utf-8 -*-
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
import os

EVID = r"F:\BIG_DATA\li16\evidencias"
OUT = r"F:\BIG_DATA\LI16_Informe_EduardoAntezana.docx"

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.7)
    s.bottom_margin = Inches(0.7)
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def center(text, size=16, bold=True, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    return p


def req(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def shot(filename):
    path = os.path.join(EVID, filename)
    if not os.path.isfile(path):
        raise SystemExit("Falta la captura " + path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(6.4))
    c = doc.add_paragraph("Archivo de evidencia: li16/evidencias/" + filename)
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)
    c.paragraph_format.space_after = Pt(8)


# Portada
center("UNIVERSIDAD DEL VALLE", 20)
center("Big Data", 18)
doc.add_paragraph()
center("LI16 - Práctica con Apache Hive", 18)
doc.add_paragraph()
center("Informe de laboratorio", 14, bold=False)
doc.add_paragraph()
center("Estudiante: Eduardo Antezana", 14)
center("Código de práctica: LI16", 14)
center("Tiquipaya - Cochabamba - 2026", 12, bold=False)
doc.add_page_break()

heading("1. Repositorio y arranque del clúster")
body(
    "Usé el repositorio que ya estaba clonado en F:\\BIG_DATA\\bigdata-grupo3 "
    "(https://github.com/EduPantoja04/bigdata-grupo3.git). No volví a clonarlo ni a descargar el CSV. "
    "El archivo de Compose está en bigdata-grupo3\\hadoop-environment\\docker-compose.yml y declara "
    "la red externa hadoop_network. En este equipo Docker Compose es la versión v5.6.0, así que el "
    "comando fue docker compose y no el antiguo docker-compose. Compose avisó que el campo version: \"2.4\" "
    "está obsoleto y lo ignoró; el clúster igual levantó."
)
body(
    "Antes de docker compose up -d --build creé la red con docker network create hadoop_network y "
    "construí la imagen base hadoop-base:3.3.6, que es Hadoop 3.3.6 como indica el Dockerfile del repositorio. "
    "Los siete servicios del Compose quedaron en ejecución."
)
req("Listar los contenedores con docker ps y registrar cuáles están activos.")
shot("01_docker_ps.png")
body(
    "En esa captura se ven los siete contenedores del repositorio, todos Up (healthy). "
    "La interfaz web de HDFS queda publicada en el host en el puerto 9870: la línea del namenode muestra "
    "0.0.0.0:9870->9870/tcp. El RPC de HDFS está en 8020. El ResourceManager publica 8089 hacia el 8088 "
    "interno, el NodeManager el 8042 y el History Server el 8188. La hora del sistema, 21/10/2026, se ve en la barra de tareas."
)

heading("2. Para qué sirve cada contenedor")
body(
    "El repositorio original no trae Hive. El Compose solo define el clúster Hadoop del grupo. "
    "Para poder crear la base y las consultas de esta práctica agregué un contenedor hive-server, "
    "con Apache Hive 4.0.1, que es la versión que el sitio de Apache Hive indica como compatible con Hadoop 3.3.6. "
    "Ese contenedor se unió a la misma red hadoop_network y usa el NameNode que ya estaba corriendo. "
    "La captura siguiente es un docker ps posterior, ya con hive-server."
)
shot("06_docker_ps_hive.png")
body(
    "namenode. Es el maestro de HDFS. Guarda el espacio de nombres (directorios y archivos) y el mapa de "
    "en qué DataNode está cada bloque. No guarda el contenido de los archivos. En core-site.xml del repositorio "
    "el sistema de archivos por defecto es hdfs://namenode:8020. Publica la web en el puerto 9870 y el RPC en el 8020. "
    "Su volumen hadoop_namenode conserva los metadatos."
)
body(
    "datanode1, datanode2 y datanode3. Cada uno almacena los bloques de HDFS y le informa al NameNode, con heartbeats, "
    "que sigue vivo. En hdfs-site.xml el grupo dejó la replicación en 2, así que cada bloque busca dos copias en nodos distintos. "
    "Por eso en la web del archivo se ve Replication 2. Sus puertos 9864 no están publicados al host: se usan dentro de la red Docker."
)
body(
    "resourcemanager. Es el maestro de YARN. Recibe las aplicaciones y reparte memoria y CPU de los NodeManagers. "
    "Su interfaz web interna es el puerto 8088 y en el host se abre como localhost:8089."
)
body(
    "nodemanager1. Es el nodo que ejecuta las tareas que le asigna el ResourceManager. Su web está en el puerto 8042 del host. "
    "En esta práctica las consultas de Hive corrieron en modo local dentro del contenedor de Hive, así que no hicieron falta "
    "contenedores de YARN para leer el CSV, pero el NodeManager forma parte del clúster que levanta el repositorio."
)
body(
    "historyserver. Guarda el historial de las aplicaciones de YARN que ya terminaron. Su web está en el puerto 8188. "
    "El volumen hadoop_historyserver persiste ese historial."
)
body(
    "hive-server. Es el contenedor que agregué para esta práctica. Tiene Hive 4.0.1 sobre la misma imagen de Hadoop 3.3.6, "
    "así que habla con hdfs://namenode:8020. El metastore es Derby embebido, dentro del contenedor, y el warehouse que usa "
    "es /user/hive/warehouse en HDFS. La imagen declara los puertos 9083, 10000 y 10002; no los publiqué al host porque "
    "las sentencias se ejecutaron con docker exec y el comando hive. En Hive 4 ese comando, tal como viene, abre Beeline "
    "y espera un HiveServer2. Para seguir el enunciado, que pide entrar con hive y lanzar SQL, dejé activo el CLI clásico "
    "de Hive (USE_BEELINE_FOR_HIVE_CLI=false). Las consultas pequeñas se ejecutaron con MapReduce local, leyendo los datos en HDFS."
)

heading("3. Contenido inicial de /user/hive/warehouse/")
req("Acceder al namenode y listar /user/hive/warehouse/.")
shot("02_hdfs_warehouse_inicial.png")
body(
    "Al entrar al namenode, hdfs dfs -ls /user/hive/warehouse/ respondió que ese directorio no existe. "
    "No es un fallo del NameNode: hdfs dfs -ls / sí funcionó y mostró solo /rmstate, que es la carpeta donde YARN "
    "guarda el estado del ResourceManager. Este repositorio es un clúster Hadoop y no crea solo las carpetas de Hive. "
    "Por eso el warehouse aparece recién cuando se prepara el directorio para subir el CSV."
)

heading("4. Interfaz web de HDFS")
req("Abrir la interfaz web del NameNode y revisar /user/hive/warehouse/.")
body(
    "En docker ps el mapeo del namenode es 0.0.0.0:9870->9870/tcp. El puerto de la izquierda es el del host y el de la "
    "derecha es el del contenedor. Como ambos son 9870, abrí http://localhost:9870 y, en Utilities, el explorador de archivos "
    "en la ruta /user/hive/warehouse."
)
shot("03_hdfs_web_inicial.png")
body(
    "La página cargó en localhost:9870/explorer.html#/user/hive/warehouse y mostró el aviso de que la ruta no existe en HDFS. "
    "Coincide con el ls del punto anterior. WebHDFS está habilitado en hdfs-site.xml; el mensaje salió porque la carpeta todavía no estaba creada."
)

heading("5. CSV y carga a HDFS")
body(
    "El archivo original, con encabezado, quedó intacto en "
    "F:\\BIG_DATA\\pseudo_facebook_EduardoAntezana.csv\\pseudo_facebook_EduardoAntezana.csv. "
    "Tiene 99 003 filas y 15 columnas: userid, age, dob_day, dob_year, dob_month, gender, tenure, friend_count, "
    "friendships_initiated, likes, likes_received, mobile_likes, mobile_likes_received, www_likes y www_likes_received. "
    "La copia sin encabezado, que es la que se cargó, está en F:\\BIG_DATA\\li16\\pseudo_facebook_EduardoAntezana.csv. "
    "Le quité solo la primera línea para que Hive no intente leer los nombres como si fueran un usuario."
)
body(
    "Los tipos que usé son INT para todas las columnas numéricas, porque los valores del archivo caben en un entero "
    "(el userid máximo es 2 193 542 y likes_received llega a 261 197), y STRING para gender. "
    "Hay 2 filas con tenure en texto NA y 175 con gender NA. En una columna INT, Hive deja el NA como nulo y no bota la fila."
)
req("Copiar pseudo_facebook_EduardoAntezana.csv al namenode, subirlo a HDFS y verificar /user/hive/warehouse/.")
shot("04_copiar_csv_hdfs.png")
body(
    "Como el directorio no existía, primero ejecuté hdfs dfs -mkdir -p /user/hive/warehouse. "
    "Después docker cp copió 5,22 MB al /tmp del namenode y hdfs dfs -put lo subió al warehouse. "
    "El ls final muestra el archivo, de 5 216 842 bytes, con replicación 2. El código de salida del put fue 0."
)
req("Revisar en la interfaz web que el CSV aparezca en /user/hive/warehouse/.")
shot("05_hdfs_web_csv.png")
body(
    "En el explorador, la misma ruta ya lista pseudo_facebook_EduardoAntezana.csv, dueño root, grupo supergroup, "
    "tamaño 4,98 MB, replicación 2 y bloque de 128 MB. Es un solo bloque porque el archivo es mucho más chico que 128 MB."
)

heading("6. Base de datos y tabla")
req("Crear la base de datos db_eduardoantezana.")
shot("07_create_database.png")
body(
    "El comando fue CREATE DATABASE db_eduardoantezana. SHOW DATABASES devolvió default y db_eduardoantezana. "
    "No usé los nombres de ejemplo del enunciado (db_pepe, fb_pepe, columnas id, day, friends) porque no coinciden "
    "con el archivo ni con el nombre pedido para esta práctica."
)
req("Crear la tabla fb_eduardoantezana con las columnas reales del CSV.")
shot("08_create_table.png")
body(
    "La tabla se creó dentro de db_eduardoantezana con las 15 columnas del CSV, separadas por coma y almacenadas como texto. "
    "Hive respondió Time taken: 1,284 seconds, sin error de sintaxis."
)
req("Verificar con DESCRIBE FORMATTED que la ubicación llegue hasta la tabla fb_eduardoantezana.")
shot("09_describe_formatted.png")
body(
    "DESCRIBE FORMATTED muestra Location "
    "hdfs://namenode:8020/user/hive/warehouse/db_eduardoantezana.db/fb_eduardoantezana "
    "y el separador de campos es la coma. Hive 4 marcó la tabla como EXTERNAL_TABLE y TRANSLATED_TO_EXTERNAL = TRUE. "
    "Eso lo hace esta versión cuando la tabla no es transaccional: la traduce a externa, pero la ruta sigue siendo la del warehouse "
    "y termina en fb_eduardoantezana, que es lo que pide la práctica. El SerDe es LazySimpleSerDe y el formato de entrada es texto."
)
req("Cargar el CSV desde HDFS y mostrar registros con una consulta limitada.")
shot("10_load_y_limit.png")
body(
    "LOAD DATA INPATH movió el CSV desde /user/hive/warehouse/ hacia la carpeta de la tabla. "
    "SELECT * LIMIT 10 devolvió 10 filas. La primera es el usuario 2094382, de 14 años, hombre, con tenure 266 y el resto de contadores en cero. "
    "Esa fila es la primera del CSV sin encabezado, así que la carga no se comió los datos ni se tragó el encabezado. "
    "Fetched: 10 row(s)."
)

heading("7. Consultas")
body(
    "Cada consulta se ejecutó con hive sobre db_eduardoantezana.fb_eduardoantezana. "
    "Hive las corrió con MapReduce local y leyó el archivo en HDFS. Los números de abajo son los de esas ejecuciones."
)

req("1. Determinar el número total de usuarios del dataset.")
shot("11_total_usuarios.png")
body(
    "COUNT(*) dio 99 003 usuarios. Es el mismo número de filas de la copia sin encabezado: cada fila es una persona "
    "y el encabezado no entró a la tabla. Con eso compruebo que la carga quedó completa."
)

req("2. Determinar el número de usuarios de Facebook por encima de los 25 años.")
shot("12_mayores_25.png")
body(
    "Con WHERE age > 25 el resultado es 56 676 usuarios. El resto, 99 003 menos 56 676, son 42 327 personas de 25 años o menos. "
    "En este archivo la edad mínima es 13, así que ese resto coincide con el grupo de 13 a 25 años que uso más adelante."
)

req("3. Determinar si los usuarios masculinos o femeninos tienden a tener más amigos.")
shot("13_amigos_genero.png")
body(
    "Comparé el promedio de friend_count, no el total, porque hay más hombres que mujeres y un total sumado no diría quién tiene más amigos por persona. "
    "Las mujeres (40 254) promedian 241,97 amigos y los hombres (58 574) promedian 165,04. "
    "En este dataset las mujeres tienden a tener más amigos. Dejé fuera 175 filas cuyo gender es NA, para no mezclarlas con male o female."
)

req("4. Determinar cuántos me gusta reciben los jóvenes (13-25 años) en comparación con los usuarios mayores de 35 años.")
shot("14_likes_edad.png")
body(
    "Usé likes_received, que son los me gusta recibidos, no los que la persona dio. "
    "El grupo de 13 a 25 años tiene 42 327 usuarios y suma 8 477 550 me gusta recibidos, con un promedio de 200,29. "
    "Los mayores de 35 años son 38 037 usuarios, suman 3 982 359 y promedian 104,70. "
    "Los jóvenes reciben más me gusta en total y también por persona. "
    "Quienes tienen entre 26 y 35 años no entran en esta comparación, porque el enunciado pide 13-25 frente a mayores de 35."
)

req("5. Encontrar el número de usuarios de Facebook para cada mes de cumpleaños.")
shot("15_mes_cumpleanos.png")
body(
    "Agrupé por dob_month. Enero tiene 11 772 usuarios y es el mes con más cumpleaños en el archivo. "
    "Noviembre tiene 7 205 y es el mes con menos. Febrero 7 632, marzo 8 110, abril 7 810, mayo 8 271, junio 7 607, "
    "julio 8 021, agosto 8 266, septiembre 7 939, octubre 8 476 y diciembre 7 894. "
    "La suma de los doce meses es 99 003, o sea que ningún registro se quedó sin mes."
)

req("6. ¿Los jóvenes (13-25 años) usan más el celular o la computadora para usar Facebook?")
shot("16_jovenes_dispositivo.png")
body(
    "El CSV no trae una columna de dispositivo ni de tiempo de uso. Lo más cercano son los me gusta hechos desde el celular "
    "(mobile_likes) y los hechos desde la web (www_likes), que tomé como uso de computadora. "
    "Entre los 13 y 25 años, mobile_likes suma 5 248 117 y www_likes suma 2 349 153. "
    "Además, 24 300 jóvenes tienen más me gusta desde el celular que desde la web, 7 816 al revés y 10 211 empatan. "
    "Con esa medida, los jóvenes usan más el celular."
)

req("7. ¿Los adultos (mayores de 25 años) usan más el celular o la computadora para usar Facebook?")
shot("17_adultos_dispositivo.png")
body(
    "Apliqué la misma lectura, ahora con age > 25. Los me gusta desde el celular suman 5 257 715 y los de la web 2 597 277. "
    "29 547 adultos tienen más actividad de me gusta en el celular, 14 401 en la computadora y 12 728 empatan. "
    "También en los mayores de 25 años predomina el celular. La diferencia es un poco menos marcada que en los jóvenes, "
    "pero el celular sigue arriba tanto en la suma como en la cantidad de personas."
)

heading("8. Cierre")
body(
    "El clúster del repositorio quedó levantado, con los siete contenedores de Hadoop en estado healthy y hive-server en ejecución. "
    "No borré contenedores, volúmenes ni el CSV original. Las capturas de esta carpeta son las de la ejecución de esta sesión "
    "y en cada una se ve la hora de la barra de tareas de Windows."
)

doc.save(OUT)
print("INFORME", OUT)
print("imagenes", sum(1 for p in doc.inline_shapes))
