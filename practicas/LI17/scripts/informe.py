# -*- coding: utf-8 -*-
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
import os

EVID = r"F:\BIG_DATA\li17\evidencias"
OUT = r"F:\BIG_DATA\LI17_Informe_EduardoAntezana.docx"

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.7)
    s.bottom_margin = Inches(0.7)
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def center(text, size=16, bold=True, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(2)
    return p


def req(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def shot(filename):
    path = os.path.join(EVID, filename)
    if not os.path.isfile(path):
        raise SystemExit("Falta la captura " + path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(6.35))
    c = doc.add_paragraph("Archivo de evidencia: li17/evidencias/" + filename)
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)
    c.paragraph_format.space_after = Pt(8)


center("UNIVERSIDAD DEL VALLE", 20)
center("Big Data", 18)
doc.add_paragraph()
center("LI17 - 2.ª Práctica con Apache Hive", 18)
doc.add_paragraph()
center("Informe de laboratorio", 14, bold=False)
doc.add_paragraph()
center("Estudiante: Eduardo Antezana", 14)
center("Código de práctica: LI17", 14)
center("Tiquipaya - Cochabamba - 2026", 12, bold=False)
doc.add_page_break()

heading("1. Entorno Hive que ya estaba en marcha")
body(
    "La práctica pide entrar a Hive con el contenedor docker-hive_hive-server_1. "
    "En este equipo ese nombre no existe. El clúster de la práctica anterior sigue siendo el de "
    "bigdata-grupo3, y Hive es el contenedor hive-server (imagen li16-hive:4.0.1). "
    "HDFS se consulta por el contenedor namenode. No recreé ni eliminé contenedores, volúmenes, "
    "la base db_eduardoantezana ni sus archivos."
)
req("Acceder a Hive desde el contenedor que realmente está activo y comprobar que el clúster anterior sigue presente.")
shot("01_docker_ps.png")
body(
    "docker ps muestra hive-server, namenode, los tres datanodes, resourcemanager, nodemanager1 e historyserver en ejecución. "
    "La hora y la fecha de Windows se ven en la barra de tareas."
)
shot("04_show_databases.png")
body(
    "CREATE DATABASE IF NOT EXISTS dbec_eduardoantezana dejó la base nueva junto a db_eduardoantezana y default. "
    "La base de la práctica anterior no se tocó."
)

heading("2. Dataset y copia de trabajo")
body(
    "El ZIP original está en C:\\Users\\pc\\Downloads\\data.csv.zip (7.548.686 bytes, 21/10/2026 22:07). "
    "No lo volví a descargar ni lo modifiqué. Lo extraje a F:\\BIG_DATA\\li17\\fuente\\data.csv "
    "(45.580.638 bytes). El archivo del proyecto F:\\BIG_DATA\\data_EduardoAntezana.csv\\data_EduardoAntezana.csv "
    "queda intacto: tiene 3.080.192 bytes y solo una parte del dataset, así que no lo usé para cargar Hive."
)
body(
    "El CSV original está en Windows-1252: contiene el símbolo de libra y no es UTF-8. "
    "La copia de trabajo F:\\BIG_DATA\\li17\\data_EduardoAntezana.csv es ese mismo contenido pasado a UTF-8, "
    "con 541.909 filas de datos más el encabezado. El lector de Python csv respetó las comillas. "
    "Un corte ingenuo por comas desplaza 4.796 filas, porque la descripción trae comas internas. "
    "También hay 1.239 descripciones con comillas internas, 135.080 CustomerID vacíos, "
    "9.288 facturas que empiezan con C y 3 con A (ajustes de deuda). "
    "Quantity cabe en un entero (de -80.995 a 80.995). UnitPrice llega a tres decimales en cuatro filas, "
    "por eso el tipo final es DECIMAL(12,3) y no un decimal de dos posiciones. "
    "Todas las fechas calzan con el patrón M/d/yyyy H:mm, sin segundos y con uno o dos dígitos en mes, día y hora."
)
req("Localizar el dataset ya descargado, conservar el original y preparar data_EduardoAntezana.csv.")
shot("02_dataset_local.png")
body(
    "La captura muestra el ZIP sin cambios, el CSV extraído, la copia de trabajo y el CSV parcial que no se usó. "
    "El encabezado es InvoiceNo, StockCode, Description, Quantity, InvoiceDate, UnitPrice, CustomerID, Country. "
    "La línea de AIRLINE LOUNGE,METAL SIGN va entre comillas, así que la coma pertenece a la descripción."
)
req("Subir una copia a una ruta HDFS exclusiva de LI17, sin borrar el warehouse anterior.")
shot("03_hdfs_li17.png")
body(
    "El archivo quedó en /li17/eduardoantezana/data_EduardoAntezana.csv (43,0 MB). "
    "Las primeras líneas leídas desde HDFS coinciden con la copia de trabajo, y el encabezado de esa copia se ve en la captura anterior. "
    "El warehouse de la práctica anterior, /user/hive/warehouse/db_eduardoantezana.db, sigue listado junto a la base nueva."
)

heading("3. Por qué no sirve cortar el CSV solo con comas")
body(
    "InvoiceNo no puede ser INT: hay valores C536379 y A563185. StockCode tampoco es numérico "
    "(85123A, POST, D). El ejemplo grupoa.tablita(id, age) no describe este archivo."
)
req("Demostrar que FIELDS TERMINATED BY ',' desplaza columnas y leer el CSV con un SerDe que respete las comillas.")
shot("05_comas_vs_serde.png")
body(
    "Sobre la misma muestra, la tabla prueba_coma parte por comas. En 536381 la descripción queda como "
    "\"AIRLINE LOUNGE, la cantidad pasa a ser METAL SIGN y United Kingdom sale de la fila. "
    "La tabla prueba_csv usa org.apache.hadoop.hive.serde2.OpenCSVSerde con separador coma, "
    "comilla doble y el mismo carácter como escape, que es lo que exige el CSV de Excel (comillas duplicadas). "
    "Ahí la descripción es AIRLINE LOUNGE,METAL SIGN, la cantidad es 2, el precio 2.1, el cliente 15311 "
    "y el país United Kingdom. La fila con comilla interna, RECORD FRAME 7\" SINGLE SIZE, también queda en una sola columna. "
    "El encabezado no entra como fila: cada consulta devolvió 5 registros."
)

heading("4. Base, tabla tipada y carga")
body(
    "La tabla de aterrizaje tec_eduardoantezana_landing es externa, apunta a /li17/eduardoantezana "
    "y usa OpenCSVSerde con skip.header.line.count=1. Todas sus columnas son STRING, porque ese SerDe no aplica tipos. "
    "La tabla final tec_eduardoantezana se llena con INSERT y estos tipos, confirmados con el archivo: "
    "invoiceno STRING, stockcode STRING, description STRING, quantity INT, invoicedate TIMESTAMP, "
    "unitprice DECIMAL(12,3), customerid INT y country STRING. "
    "La fecha se convierte con unix_timestamp(texto, 'M/d/yyyy H:mm') y from_unixtime, y después a TIMESTAMP. "
    "El CustomerID vacío se vuelve NULL con NULLIF antes del CAST. "
    "La tabla final se guardó como ORC. Si se hubiera vuelto a grabar como texto separado por comas, "
    "las descripciones con coma se habrían desplazado otra vez al leerlas."
)
req("Crear la base dbec_eduardoantezana y la tabla tec_eduardoantezana, y cargar el CSV sin tomar el encabezado como registro.")
shot("06_crear_y_cargar.png")
body(
    "La captura muestra el DDL de la tabla de aterrizaje, el de la tabla tipada y el INSERT. "
    "El filtro invoiceno <> 'InvoiceNo' descarta el encabezado si el skip no lo hubiera hecho. "
    "La ruta de HDFS leída en ese trabajo es la copia de LI17."
)
req("Revisar con DESCRIBE FORMATTED que la ubicación corresponda a la tabla tec_eduardoantezana.")
shot("07_describe_formatted.png")
body(
    "Location es hdfs://namenode:8020/user/hive/warehouse/dbec_eduardoantezana.db/tec_eduardoantezana. "
    "Esa ruta es la de la tabla, no la carpeta del CSV. Hive 4 marca la tabla como EXTERNAL_TABLE "
    "con TRANSLATED_TO_EXTERNAL y external.table.purge=TRUE: es la forma en que esta versión registra "
    "una tabla creada en el warehouse. El SerDe es OrcSerde. Después de ANALYZE TABLE, numRows es 541909 "
    "y totalSize es 2784817, el tamaño del ORC generado en la carga."
)
req("Verificar conteos, nulos, una fila con coma interna y los primeros 10 registros.")
shot("08_conteo_y_alineacion.png")
shot("08b_diez_filas.png")
body(
    "Hay 0 filas con el texto InvoiceNo, así que el encabezado no se cargó. "
    "Las fechas nulas son 0. Los CustomerID nulos son 135.080, la misma cantidad de celdas vacías del CSV. "
    "Las facturas que empiezan con C son 9.288. "
    "La fila 536381 / 82567 queda alineada: descripción AIRLINE LOUNGE,METAL SIGN, cantidad 2, "
    "fecha 2010-12-01 09:41:00, precio 2.100, cliente 15311 y país United Kingdom. "
    "Las diez primeras filas tienen las ocho columnas en orden y las fechas ya como timestamp. "
    "DESCRIBE FORMATTED reporta las 541.909 filas de datos del archivo."
)

heading("5. Conversión de InvoiceDate")
body(
    "En el CSV la fecha es texto con el patrón M/d/yyyy H:mm. No trae segundos. "
    "El mes, el día y la hora ocupan uno o dos dígitos: 12/1/2010 8:26 y 8/12/2011 14:50 son los dos casos. "
    "La máscara de Hive es exactamente 'M/d/yyyy H:mm'. YEAR, MONTH, HOUR y DAYOFWEEK se aplican al TIMESTAMP ya convertido, "
    "no a la cadena original. En Hive DAYOFWEEK usa 1 para domingo y 7 para sábado."
)
req("Probar, antes de las consultas temporales, que InvoiceDate se convierte a fecha y hora.")
shot("09_conversion_fechas.png")
body(
    "536365, cuyo texto era 12/1/2010 8:26, queda 2010-12-01 08:26:00: año 2010, mes 12, día 1, hora 8, minuto 26 "
    "y día de semana 4 (miércoles). A563185, cuyo texto era 8/12/2011 14:50, queda 2011-08-12 14:50:00: "
    "hora 14 y día de semana 6 (viernes). Las fechas no parseadas son 0."
)

heading("6. Nivel básico")

req("¿Cuál es el ingreso total generado por ventas válidas?")
shot("10_q01_ingreso_total.png")
body(
    "La consulta suma Quantity * UnitPrice, redondea a dos decimales y se queda con cantidad positiva, "
    "precio positivo y CustomerID no nulo. El resultado es 8911407.90. "
    "Responde al ingreso de esas ventas válidas. Ese total deja fuera las filas sin CustomerID: "
    "más adelante esas filas suman 1755276.64, así que 8911407.90 no representa todo el archivo."
)

req("¿Cuáles son los 10 productos más vendidos por cantidad?")
shot("11_q02_top_productos.png")
body(
    "Agrupa por Description, suma Quantity cuando es positiva y cuenta facturas distintas. "
    "Ordena por unidades y corta en 10. El primero es PAPER CRAFT , LITTLE BIRDIE, con 80.995 unidades en un solo pedido: "
    "es el valor máximo del archivo y concentra el ranking en un outlier. Siguen MEDIUM CERAMIC TOP STORAGE JAR "
    "(78.033 unidades, 247 pedidos), WORLD WAR 2 GLIDERS ASSTD DESIGNS (55.047, 535), "
    "JUMBO BAG RED RETROSPOT (48.478, 2.092) y WHITE HANGING HEART T-LIGHT HOLDER (37.895, 2.260). "
    "La novena fila es NULL, con 32.547 unidades en 592 pedidos: son descripciones vacías convertidas a nulo. "
    "La coma de PAPER CRAFT , LITTLE BIRDIE sigue dentro de la descripción, así que las columnas no se corrieron. "
    "La consulta responde el ranking literal por unidades; ese ranking está dominado por cantidades extremas de pocos pedidos."
)

req("¿Cuántos clientes únicos hay por país?")
shot("12_q03_clientes_pais.png")
body(
    "Cuenta CustomerID distintos por país y calcula el porcentaje sobre el total de esa cuenta, "
    "solo en filas con CustomerID. Devolvió 37 países. Reino Unido concentra 3.950 clientes, el 90,18 %. "
    "Siguen Alemania con 95 (2,17 %) y Francia con 87 (1,99 %). "
    "La captura muestra el final del orden descendente, hasta los países con un cliente (0,02 %), y el SQL completo. "
    "Responde la distribución de clientes identificados. No incluye a quienes compraron sin CustomerID."
)

req("¿Cuál es el ticket promedio por orden?")
shot("13_q04_ticket.png")
body(
    "Primero arma el total de cada InvoiceNo con cantidad y precio positivos, y después calcula el promedio "
    "y la mediana aproximada de esos totales. El ticket promedio es 534.40 y la mediana es 303.79922222222217. "
    "El promedio queda por encima de la mediana, así que unos pocos pedidos altos jalan la media. "
    "Esta consulta sí está armada por orden. A diferencia de la del ingreso total, aquí no se exige CustomerID, "
    "así que entran pedidos con y sin cliente identificado. PERCENTILE_APPROX es la mediana aproximada de Hive, "
    "no un percentil exacto."
)

req("¿Cómo evolucionan las ventas mes a mes?")
shot("14_q05_evolucion_mensual.png")
body(
    "Usa YEAR y MONTH sobre el timestamp ya convertido, cuenta facturas distintas y suma el importe de líneas "
    "con cantidad y precio positivos. Hay 13 meses, de diciembre de 2010 a diciembre de 2011. "
    "El punto más alto es noviembre de 2011: 2.769 pedidos y 1509496.33 de ingresos. "
    "Diciembre de 2011 baja a 819 pedidos y 638792.68. Octubre de 2011 es el segundo mes en ingresos (1154979.30). "
    "La consulta responde la serie mensual pedida. Diciembre de 2011 no se puede leer como un mes comparable "
    "al resto solo con esta tabla, porque el resultado mismo ya lo muestra mucho más bajo en pedidos."
)

heading("7. Nivel intermedio")

req("¿Qué productos tienen mayor tasa de devolución o cancelación?")
shot("15_q06_devolucion.png")
body(
    "La tasa es de líneas, no de pedidos. Para cada descripción cuenta las líneas con cantidad negativa o factura que empieza con C, "
    "las compara con las líneas de cantidad positiva y divide entre todas las líneas del producto. "
    "Solo entran filas con CustomerID y productos con al menos 10 líneas. "
    "CRUK Commission queda en 100 % (16 líneas de devolución y 0 ventas) y Discount también en 100 % (77 y 0): "
    "son ajustes, no mercadería devuelta. El primer artículo con ventas positivas en este corte es "
    "WOODEN BOX ADVENT CALENDAR, con 8 devoluciones, 5 ventas y 61,54 %. "
    "La consulta responde el ranking que define su propio SQL. No es el porcentaje de pedidos cancelados ni de unidades devueltas."
)

req("¿En qué hora del día hay más ventas y mayor ticket?")
shot("16_q07_hora.png")
body(
    "Agrupa por HOUR del timestamp. La hora 12 concentra más pedidos: 3.220. "
    "La hora 10 concentra más ingresos: 1446742.70, apenas por encima de las 12 (1444245.88). "
    "La columna ticket_promedio es el promedio de Quantity * UnitPrice por línea, no el ticket de la orden. "
    "Por eso las 7 de la mañana marcan 81.95 con solo 29 pedidos, y las 6 marcan 4.25 con un pedido. "
    "La consulta responde pedidos, ingresos y promedio por línea en cada hora. "
    "El mayor volumen de pedidos está a las 12 y el mayor ingreso a las 10; el promedio por línea no señala esa misma hora."
)

req("¿Qué productos son estacionales? Ventas por mes.")
shot("17_q08_estacionalidad.png")
body(
    "Esta es la consulta completa del enunciado. Para cada descripción y mes suma las unidades positivas "
    "y calcula qué porcentaje representan dentro del total de esa descripción. "
    "Se queda con los grupos de más de 50 unidades. Hive devolvió 14.927 filas. "
    "La captura muestra el final de ese resultado, el conteo Fetched: 14927 row(s) y el SQL entero. "
    "La salida completa quedó en F:\\BIG_DATA\\li17\\salidas\\q08_estacionalidad.log. "
    "Al final aparecen descripciones de error (wrongly coded, incorrectly credited) y la descripción NULL repartida en varios meses. "
    "El porcentaje por mes responde la distribución pedida. Un mes al 100 % solo dice que, dentro del filtro de más de 50 unidades, "
    "esas unidades cayeron en un único mes; no demuestra por sí solo un calendario comercial."
)
req("Complemento de la estacionalidad, para leer el extremo del resultado. No sustituye la consulta anterior.")
shot("18_q08b_complemento.png")
body(
    "El complemento repite la misma agregación y el mismo HAVING, y además exige distribucion_pct mayor o igual a 40, "
    "ordenado por ese porcentaje y por unidades, con límite 15. Las 15 filas visibles están al 100 % en un solo mes. "
    "La primera es otra vez PAPER CRAFT , LITTLE BIRDIE: 80.995 unidades en el mes 12. "
    "También entran textos de corrección de datos, como incorrectly credited y wrongly coded. "
    "Sirve para mostrar el extremo; la evidencia de la pregunta sigue siendo la consulta de 14.927 filas."
)

req("¿Qué productos tienen mayor margen implícito? El precio unitario se usa como proxy.")
shot("19_q09_margen_proxy.png")
body(
    "Calcula el precio promedio, los pedidos distintos y el producto de ambos, con cantidad y precio positivos "
    "y al menos cinco pedidos. El primer score es DOTCOM POSTAGE: precio promedio 292.14, 706 pedidos y score 206248.77. "
    "Siguen Manual (66743.93) y POSTAGE (34992.23). El primer artículo de catálogo en el corte es "
    "REGENCY CAKESTAND 3 TIER, con precio 13.98, 1.988 pedidos y score 27786.13. "
    "El score multiplica precio por rotación de pedidos. El dataset no trae costo de producto, "
    "así que este número es un proxy de precio por rotación y no un margen real."
)

req("¿Hay diferencia de comportamiento entre clientes con identificador y sin identificador?")
shot("20_q10_clientes_id.png")
body(
    "Separa las líneas de cantidad y precio positivos según CustomerID nulo o no. "
    "Anónimo: 1.428 pedidos, promedio por línea 13.28 e ingreso 1755276.64. "
    "Registrado: 18.532 pedidos, promedio por línea 22.40 e ingreso 8911407.90. "
    "El ingreso registrado coincide con la primera consulta, porque esa consulta también exigía CustomerID. "
    "El promedio de esta consulta es por línea, no por pedido. "
    "Los clientes identificados concentran más pedidos y más ingreso. "
    "Si un reporte usa solo las ventas con CustomerID, deja fuera los 1755276.64 y los 1.428 pedidos anónimos."
)

req("¿Qué días de la semana tienen mayor conversión?")
shot("21_q11_dia_semana.png")
body(
    "La consulta agrupa por DAYOFWEEK del timestamp, pone el nombre del día y suma pedidos e ingresos "
    "de líneas con cantidad y precio positivos. Devolvió 6 filas: domingo a viernes. "
    "No aparece el sábado (día 7). El jueves tiene la mayor cantidad de pedidos, 4.246, y el mayor ingreso, 2203161.24. "
    "El martes es el segundo ingreso (2178632.61) con 3.554 pedidos. El domingo es el más bajo: 2.204 pedidos y 813827.61. "
    "El dataset no contiene visitas ni sesiones. La consulta compara pedidos e ingresos por día; "
    "no calcula una tasa de conversión."
)

req("¿Cuál es la tasa de cancelación por país?")
shot("22_q12_cancelacion_pais.png")
body(
    "La tasa también es de líneas. Cuenta líneas con cantidad negativa o factura que empieza con C, "
    "las ventas con cantidad positiva y el porcentaje sobre todas las líneas del país. "
    "Exige al menos 50 líneas y no filtra CustomerID. Devolvió 32 países. "
    "Estados Unidos queda primero, con 38,49 % (112 líneas de cancelación y 179 de venta). "
    "Malta tiene 11,81 % y Japón 10,34 %. Reino Unido, con muchísimas más líneas, queda en 1,86 % "
    "(9.192 cancelaciones y 486.286 ventas). "
    "La consulta responde esa tasa por línea y con ese filtro. No es la proporción de pedidos cancelados "
    "ni una tasa calculada solo sobre clientes identificados. En países con pocas líneas el porcentaje se mueve con pocos registros."
)

heading("8. Cierre")
body(
    "La base dbec_eduardoantezana y la tabla tec_eduardoantezana quedaron cargadas con las 541.909 filas del CSV completo. "
    "El ZIP original, el CSV parcial del proyecto y la base db_eduardoantezana siguen en su sitio. "
    "Los contenedores hive-server y namenode no se recrearon. "
    "Las doce consultas del enunciado se ejecutaron sobre dbec_eduardoantezana.tec_eduardoantezana, "
    "y la de estacionalidad se conservó completa además del complemento de 15 filas."
)

doc.save(OUT)
print(OUT)
