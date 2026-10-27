from docx import Document
from pathlib import Path
from hashlib import sha256
p=Path(r'F:\BIG_DATA\LI17_Informe_EduardoAntezana.docx'); d=Document(p); folder=Path(r'F:\BIG_DATA\li17\evidencias')
for i,para in enumerate(d.paragraphs):
    rels=para._p.xpath('.//a:blip/@r:embed')
    if rels:
        prev=' | '.join(d.paragraphs[j].text for j in range(max(0,i-2),i) if d.paragraphs[j].text.strip())
        for rel in rels:
            part=d.part.related_parts[rel]; data=part.blob
            print(i,prev[-95:],part.partname, len(data),sha256(data).hexdigest()[:12])
print('HIGHLIGHTS')
for i,para in enumerate(d.paragraphs):
    if any(r.font.highlight_color for r in para.runs): print(i,para.text)
